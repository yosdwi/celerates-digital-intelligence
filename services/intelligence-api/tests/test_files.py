"""Company Files (doc 17, ADR-018): extraction, registry, class/record authorization, model visibility, ERP sync,
ingestion queue, Agent tools and HTTP surface. ERP is faked at its contract boundary (access, readable, feed, bytes)."""

import hashlib
import io
import json
import types
from uuid import uuid4

import pytest
from cryptography.hazmat.primitives.serialization import Encoding, PublicFormat
from fastapi.testclient import TestClient
from test_agent import KEY, SECOND, USER, AskingERP, docx_bytes, events, mint

from cdi import delegation, extract, files
from cdi.agent import playbooks
from cdi.api import app
from cdi.config import settings
from cdi.db import all_rows, connect, one
from cdi.migrate import migrate

PMO_USER = "7c2d4e5f-3a4b-4c5d-8e6f-7a8b9c0d1e2f"


@pytest.fixture(scope="module", autouse=True)
def schema():
    migrate()


@pytest.fixture(autouse=True)
def keys(monkeypatch):
    pem = KEY.public_key().public_bytes(Encoding.PEM, PublicFormat.SubjectPublicKeyInfo).decode()
    monkeypatch.setattr(settings(), "erp_delegation_public_keys", json.dumps({"k1": pem}))
    monkeypatch.setattr(settings(), "generation_mode", "demo")
    monkeypatch.setattr(settings(), "embedding_mode", "demo")
    with connect() as conn:  # each test starts from an empty registry
        conn.execute("DELETE FROM files")
        conn.execute("DELETE FROM file_access_log")


def pdf(pages):
    """A PDF with one page per entry: a list of text lines, or [] for a page with no text layer (a scan)."""
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", None, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for lines in pages:
        content = (
            "BT /F1 12 Tf 72 720 Td " + " ".join(f"({line}) Tj 0 -16 Td" for line in lines) + " ET" if lines else ""
        )
        objs.append(f"<< /Length {len(content)} >>\nstream\n{content}\nendstream")
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {len(objs)} 0 R "
            "/Resources << /Font << /F1 3 0 R >> >> >>"
        )
        kids.append(f"{len(objs)} 0 R")
    objs[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"
    out, offsets = "%PDF-1.4\n", []
    for i, obj in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{obj}\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n" + "".join(f"{o:010d} 00000 n \n" for o in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF"
    return out.encode("latin-1")


def xlsx(rows, title="Kebutuhan"):
    from openpyxl import Workbook

    book = Workbook()
    ws = book.active
    ws.title = title
    for row in rows:
        ws.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


class FakeDocling:
    """Stands in for Docling's converter: items with page provenance, a table, OCR text for the scanned page."""

    def __init__(self, pages):
        self.pages, self.calls = pages, []

    def convert(self, path):
        self.calls.append(path)
        items = []
        for page, entries in self.pages.items():
            for kind, text in entries:
                prov = [types.SimpleNamespace(page_no=page)]
                if kind == "table":
                    items.append(
                        types.SimpleNamespace(label="table", prov=prov, export_to_markdown=lambda doc=None, t=text: t)
                    )
                else:
                    items.append(types.SimpleNamespace(label=kind, prov=prov, text=text))
        document = types.SimpleNamespace(iterate_items=lambda: [(i, 0) for i in items])
        return types.SimpleNamespace(document=document)


# ── extraction ─────────────────────────────────────────────────────────────────────────────────────────────────────
def test_extraction_keeps_pages_tables_and_ocr_and_detects_identity_patterns():
    with pytest.raises(extract.ExtractError):
        extract.sniff("kontrak.pdf", b"MZ\x90\x00 not a pdf")
    assert extract.sniff("sop.PDF", pdf([["SOP"]])) == (".pdf", "application/pdf")

    got = extract.extract(
        "sop.pdf", pdf([["SOP Rekrutmen", "Setiap requisition wajib punya TA PIC."], ["Lampiran", "Formulir A"]])
    )
    assert got.parser == "pypdf-v1" and [p for p, _ in got.pages] == [1, 2]
    pieces = extract.chunks(got.pages)
    assert {c["page"] for c in pieces} == {1, 2} and all(c["page"] in (1, 2) for c in pieces)

    doc = extract.extract(
        "pks.docx",
        docx_bytes(
            ["PERJANJIAN KERJA SAMA", "Pasal 1 Jangka waktu 12 bulan."], [["Posisi", "Jumlah"], ["Engineer", "8"]]
        ),
    )
    assert doc.tables == 1
    table = [c for c in extract.chunks(doc.pages) if c["block"] == "table"]
    assert table and "Engineer | 8" in table[0]["text"]

    sheet = extract.extract(
        "manpower.xlsx", xlsx([["Posisi", "Jumlah", "Level"], ["Backend Engineer", 3, "Senior"], ["QA", 1, "Junior"]])
    )
    assert sheet.parser == "sheets-v1" and sheet.tables == 1
    assert "Posisi | Jumlah | Level\nBackend Engineer | 3 | Senior" in sheet.text

    # A scanned page goes to Docling with OCR; blocks from it are marked `ocr`; tables keep their page.
    scan = pdf([["BAST PT Synthetic", "Periode Maret 2026"], []])
    fake = FakeDocling(
        {
            1: [("title", "BAST PT Synthetic"), ("text", "Periode Maret 2026")],
            2: [("text", "Diterima oleh PT Synthetic"), ("table", "| Item | Qty |\n| Jasa | 1 |")],
        }
    )
    got = extract.extract("bast.pdf", scan, converter=fake)
    assert fake.calls and got.parser == "docling-v2+ocr" and got.ocr_pages == 1 and got.tables == 1
    kinds = {(p, b["kind"]) for p, blocks in got.pages for b in blocks}
    assert (2, "ocr") in kinds and (2, "table") in kinds and (1, "heading") in kinds and (1, "text") in kinds
    # Without OCR available (the synchronous Agent drop) a scan is flagged instead.
    blank = extract.extract("scan.pdf", pdf([[]]), allow_ocr=False)
    assert blank.scanned and not blank.pages

    assert extract.pii_flags("NIK: 3171012345678901 KARTU TANDA PENDUDUK") == ["ktp", "nik"]
    assert extract.pii_flags("NPWP 01.234.567.8-901.000") == ["npwp"]
    assert extract.pii_flags("Kontrak berlaku 12 bulan senilai 250000000") == []


# ── registry and authorization ────────────────────────────────────────────────────────────────────────────────────
class FilesERP:
    """ERP at the contract boundary: per-user class access and record readability, as ERP would decide them."""

    access = {}
    readable_refs = set()
    readable_entities = set()

    def __init__(self, principal):
        self.principal = principal

    def files_access(self):
        return FilesERP.access.get(
            self.principal.sub, {"owner": False, "general": True, "divisions": [], "commercial": [], "personal": []}
        )

    def files_readable(self, refs, entities):
        return {
            "refs": [r for r in refs if r in FilesERP.readable_refs],
            "entities": [e for e in entities if e in FilesERP.readable_entities],
        }


def who(sub, **access):
    principal = delegation.DelegatedPrincipal(delegation.verify(mint(sub=sub, owner=access.get("owner", False))), "t")
    FilesERP.access[sub] = {
        "owner": False,
        "general": True,
        "divisions": [],
        "commercial": [],
        "personal": [],
        **access,
    }
    return principal


def drain():
    while files.tick():
        pass


@pytest.fixture
def erp(monkeypatch):
    monkeypatch.setattr(files, "_erp", lambda principal: FilesERP(principal))
    FilesERP.access, FilesERP.readable_refs, FilesERP.readable_entities = {}, set(), set()
    return FilesERP


def test_classes_compose_with_erp_access_and_model_visibility(erp):
    owner = who(USER, owner=True)
    sales = who(SECOND, divisions=["sales"])
    pmo = who(PMO_USER, divisions=["pmo"], commercial=["pmo"])

    sop = files.create_managed(
        sales, "sop.pdf", pdf([["SOP Penagihan", "Invoice dikirim setiap tanggal 5."]]), kind="sop"
    )
    assert sop["access_class"] == "general" and sop["versions"][0]["ingest_state"] == "queued"
    # A kind's default class can only be tightened by the uploader; loosening needs an Owner.
    with pytest.raises(files.FilesError):
        files.create_managed(sales, "cv.pdf", pdf([["CV"]]), kind="cv", access_class="general")
    with pytest.raises(files.FilesError):  # the uploader must hold the class for that division
        files.create_managed(sales, "pks.pdf", pdf([["PKS"]]), kind="contract", owner_division="pmo")
    contract = files.create_managed(
        pmo,
        "pks-astra.pdf",
        pdf([["PERJANJIAN KERJA SAMA PT Astra", "Nilai kontrak Rp 250.000.000 selama 12 bulan"]]),
        kind="contract",
        owner_division="pmo",
        title="PKS Astra 2026",
    )
    proposal = files.create_managed(
        sales, "proposal.pdf", pdf([["Proposal Astra", "Tim 8 engineer"]]), kind="proposal", owner_division="sales"
    )
    drain()
    assert files.detail(owner, contract["id"])["versions"][0]["ingest_state"] == "indexed"

    # General: everyone with ERP access; the model sees content (policy `full`).
    hits = files.search(pmo, "penagihan invoice", purpose="agent")
    assert [h["id"] for h in hits] == [sop["id"]] and "tanggal 5" in hits[0]["snippet"] and hits[0]["page"] == 1
    # Division: only readers of the owning division.
    assert [h["id"] for h in files.search(sales, "proposal astra")] == [proposal["id"]]
    assert proposal["id"] not in [h["id"] for h in files.search(pmo, "proposal astra")]
    # Commercial: only holders of the class for that division; the model never receives the content.
    ids = lambda principal, q="perjanjian astra": [h["id"] for h in files.search(principal, q)]  # noqa: E731
    assert contract["id"] not in ids(sales)
    ui = files.search(pmo, "perjanjian astra", purpose="ui")
    agent = files.search(pmo, "perjanjian astra", purpose="agent")
    assert ui[0]["id"] == contract["id"] and "250.000.000" in ui[0]["snippet"]
    assert agent[0]["id"] == contract["id"] and "snippet" not in agent[0] and agent[0]["content_shared"] is False
    withheld = files.read(pmo, contract["id"], "nilai", purpose="agent")
    assert withheld["withheld"] is True and withheld["passages"] == [] and withheld["file"]["content_shared"] is False
    assert files.read(pmo, contract["id"], "nilai", purpose="ui")["passages"][0]["page"] == 1
    with pytest.raises(files.FilesError):
        files.read(sales, contract["id"])
    # A managed commercial file linked to a record also needs ERP to confirm the user can read that record.
    record = str(uuid4())
    erp.readable_entities = {f"project_document:{record}"}  # the Owner can read it, so may link it
    files.act(
        owner, contract["id"], "link", link={"type": "project_document", "id": record, "label": "Dokumen proyek Astra"}
    )
    assert contract["id"] in ids(pmo)
    erp.readable_entities = set()
    assert contract["id"] not in ids(pmo), "linked record not readable → hidden"
    # Owners see every class (ERP confirms Owners can read every record).
    erp.readable_entities = {f"project_document:{record}"}
    assert {h["id"] for h in files.search(owner, "astra")} >= {contract["id"], proposal["id"]}


def test_identity_patterns_hold_a_broad_file_and_withdrawal_unindexes(erp):
    owner = who(USER, owner=True)
    sales = who(SECOND, divisions=["sales"])
    pmo = who(PMO_USER, divisions=["pmo"])
    held = files.create_managed(
        sales, "data.pdf", pdf([["KARTU TANDA PENDUDUK", "NIK 3171012345678901"]]), kind="admin"
    )
    drain()
    row = files.detail(sales, held["id"])
    assert row["state"] == "pending_review" and row["versions"][0]["flags"] == ["ktp", "nik"]
    with pytest.raises(files.FilesError):
        files.detail(pmo, held["id"])  # invisible to others while held
    with pytest.raises(files.FilesError):
        files.act(sales, held["id"], "reclassify", access_class="general")  # releasing is an Owner decision
    assert files.act(owner, held["id"], "reclassify", access_class="personal", owner_division="hr")["state"] == "active"

    sop = files.create_managed(sales, "sop.txt", b"SOP cuti: ajukan tiga hari sebelumnya.", kind="sop")
    drain()
    assert files.search(pmo, "cuti")
    files.act(sales, sop["id"], "withdraw")
    assert not files.search(pmo, "cuti") and not files.search(sales, "cuti")
    with connect() as conn:
        assert one(conn, "SELECT count(*)::int AS n FROM chunks WHERE file_id=%s", (sop["id"],))["n"] == 0


def test_erp_files_sync_ingest_and_record_checks(erp, monkeypatch):
    monkeypatch.setattr(settings(), "erp_mode", "http")
    monkeypatch.setattr(files, "catalog", lambda: files.DEFAULT_CATALOG)
    candidate, invoice = str(uuid4()), str(uuid4())
    cv_key, po_key = f"attachment:{uuid4()}", f"attachment:{uuid4()}"
    cv = pdf([["CURRICULUM VITAE Budi Synthetic", "Java Spring Boot 7 tahun, pernah di PT Astra"]])
    bast = pdf([["BAST Maret 2026 PT Synthetic"], []])
    feed = {
        cv_key: {
            "kind": "cv",
            "access_class": "personal",
            "owner_division": "ta",
            "name": "cv-budi.pdf",
            "entity": {
                "type": "candidate",
                "id": candidate,
                "label": "CAND-1 · Budi Synthetic",
                "href": f"/ta/candidates/{candidate}/edit",
            },
            "etag": "k1",
            "body": cv,
        },
        f"column:project_invoices.bast_support_doc_url:{invoice}": {
            "kind": "bast",
            "access_class": "commercial",
            "owner_division": "pmo",
            "name": "BAST",
            "entity": {
                "type": "invoice",
                "id": invoice,
                "label": "Invoice · PT Synthetic",
                "href": "/pmo/invoices/x/edit",
            },
            "etag": "k2",
            "body": bast,
        },
        po_key: {
            "kind": "po",
            "access_class": "commercial",
            "owner_division": "sales",
            "name": "https://drive.example/po",
            "entity": None,
            "etag": "https://drive.example/po",
            "external": True,
        },
    }

    def machine(path, params=None, raw=False):
        if path == "feed":
            items = [
                {
                    "ref": ref,
                    "origin": "external" if v.get("external") else "erp",
                    "kind": v["kind"],
                    "access_class": v["access_class"],
                    "owner_division": v["owner_division"],
                    "name": v["name"],
                    "etag": v["etag"],
                    "open_url": f"/api/documents?path={ref}",
                    "entity": v["entity"],
                }
                for ref, v in sorted(feed.items())
            ]
            return {"items": items, "next": None}
        if path == "content":
            v = feed.get(params["ref"])
            return types.SimpleNamespace(content=v["body"]) if v and "body" in v else None
        raise AssertionError(path)

    monkeypatch.setattr(files, "_machine", machine)
    assert files.sync_erp() == {"seen": 3, "new": 3, "versions": 3, "withdrawn": 0}
    ocr = FakeDocling(
        {1: [("text", "BAST Maret 2026 PT Synthetic")], 2: [("text", "Diterima dan disetujui PT Synthetic")]}
    )
    while files.tick(converter=ocr):
        pass
    with connect() as conn:
        states = {
            r["kind"]: (r["ingest_state"], r["ocr_pages"])
            for r in all_rows(
                conn, "SELECT f.kind, v.ingest_state, v.ocr_pages FROM files f JOIN file_versions v ON v.file_id=f.id"
            )
        }
    assert states == {"cv": ("indexed", 0), "bast": ("indexed", 1), "po": ("metadata_only", None)}

    ta = who(SECOND, divisions=["ta"], personal=["ta"])
    cv_ref = next(r for r, v in feed.items() if v["kind"] == "cv")
    assert not files.search(ta, "java spring"), "ERP has not confirmed the CV is readable"
    erp.readable_refs = {cv_ref}
    hit = files.search(ta, "java spring", purpose="agent")[0]
    assert hit["kind"] == "cv" and hit["links"][0]["label"] == "CAND-1 · Budi Synthetic" and "snippet" not in hit
    assert "Java" in files.search(ta, "java spring", purpose="ui")[0]["snippet"]
    sales = who(PMO_USER, divisions=["sales"])
    assert not files.search(sales, "java spring"), "no personal grant → nothing, not even a count"
    # OCR'd text is searchable, with its page.
    pmo = who(USER, owner=True)
    bast_ref = next(r for r, v in feed.items() if v["kind"] == "bast")
    erp.readable_refs = {cv_ref, bast_ref}
    assert files.search(pmo, "diterima disetujui")[0]["page"] == 2

    # ERP deletes the attachment and replaces the BAST: withdrawn and re-versioned on the next pass.
    del feed[cv_ref]
    feed[bast_ref]["etag"] = "k3"
    assert files.sync_erp() == {"seen": 2, "new": 0, "versions": 1, "withdrawn": 1}
    erp.readable_refs = {cv_ref}
    assert not files.search(ta, "java spring")
    with connect() as conn:
        assert one(conn, "SELECT current_version FROM files WHERE origin_ref=%s", (bast_ref,))["current_version"] == 2

    # A version ERP no longer serves fails, backs off, and is retried by a curator.
    feed[bast_ref].pop("body")
    files.tick(converter=ocr)
    with connect() as conn:
        failed = one(conn, "SELECT file_id, version, attempts, error FROM file_versions WHERE ingest_state='failed'")
    assert failed["attempts"] == 1 and failed["error"].startswith("LookupError")
    assert not files.claim(), "backoff: not claimed again at once"
    assert files.retry(failed["file_id"], failed["version"])


def test_agent_finds_files_but_never_carries_model_hidden_content(erp, monkeypatch):
    class ERP(AskingERP, FilesERP):
        pass

    owner = who(USER, owner=True)
    secret = "Nilai kontrak Rp 987.654.321"
    contract = files.create_managed(
        owner, "pks.pdf", pdf([["PERJANJIAN KERJA SAMA PT Astra", secret]]), kind="contract", owner_division="pmo"
    )
    sop = files.create_managed(
        owner, "sop.pdf", pdf([["SOP Kontrak", "Setiap kontrak ditinjau legal sebelum ditandatangani."]]), kind="sop"
    )
    drain()
    token = mint(ctx={"path": "/pmo", "module": "pmo"})
    monkeypatch.setattr(settings(), "erp_mode", "http")
    monkeypatch.setattr(settings(), "erp_base_url", "http://erp.invalid/api/integration/v1")
    monkeypatch.setattr(settings(), "erp_token", "x" * 40)
    monkeypatch.setattr(settings(), "erp_action_token", "y" * 40)
    monkeypatch.setattr(settings(), "api_access_token", "w" * 40)
    monkeypatch.setattr(files, "catalog", lambda: files.DEFAULT_CATALOG)
    monkeypatch.setattr(playbooks, "DelegatedERP", ERP)
    monkeypatch.setattr(files, "_erp", lambda principal: FilesERP(principal))
    monkeypatch.setattr(playbooks, "start", lambda r, u, a: playbooks.execute(r, u, a))
    with TestClient(app) as c:

        def ask(args):
            run_id = str(uuid4())
            c.post(
                "/api/agent/runs",
                json={"run_id": run_id, "skill": "ask", "args": args},
                headers={"X-ERP-Delegation": token},
            )
            return run_id, events(c, run_id, token)

        run_id, stream = ask({"query": "cari kontrak astra"})
        assert stream[-1][1]["type"] == "RUN_FINISHED", stream[-1]
        cards = [
            i
            for _, e in stream
            if e.get("name") == "celerates.evidence"
            for i in e["value"]["items"]
            if i["type"] == "file"
        ]
        assert {c["source"]["file_id"] for c in cards} >= {contract["id"]}
        text = "".join(e["delta"] for _, e in stream if e["type"] == "TEXT_MESSAGE_CONTENT")
        assert "PKS" in text or "pks" in text.lower()
        with connect() as conn:
            steps = json.dumps(
                [r["event"] for r in all_rows(conn, "SELECT event FROM agent_steps WHERE run_id=%s", (run_id,))]
            )
        assert "987.654.321" not in steps, "commercial content never enters the run"

        # "Tanya file ini": a general file answers with cited pages; a model-hidden class is refused.
        _, stream = ask({"query": "kapan kontrak ditinjau?", "file_id": sop["id"]})
        assert stream[-1][1]["type"] == "RUN_FINISHED"
        assert any(
            i["type"] == "file" and "hal. 1" in i["title"]
            for _, e in stream
            if e.get("name") == "celerates.evidence"
            for i in e["value"]["items"]
        )
        _, stream = ask({"query": "berapa nilainya?", "file_id": contract["id"]})
        assert stream[-1][1]["type"] == "RUN_ERROR" and "tidak dibagikan" in stream[-1][1]["message"]
        assert secret not in json.dumps([e for _, e in stream])


def test_http_surface_upload_search_content_and_console(erp, monkeypatch):
    import hashlib as _h

    curator_token = "f" * 40
    monkeypatch.setattr(
        settings(),
        "intelligence_principals_json",
        json.dumps(
            [
                {
                    "id": "curator",
                    "token_sha256": _h.sha256(curator_token.encode()).hexdigest(),
                    "roles": ["reviewer", "curator"],
                    "divisions": ["sales"],
                }
            ]
        ),
    )
    monkeypatch.setattr(settings(), "api_access_token", "")
    who(USER, owner=False, divisions=["sales"])
    who(SECOND, divisions=["pmo"])
    token, other = mint(owner=False), mint(sub=SECOND, owner=False)
    body = b"SOP onboarding: laptop disiapkan H-3."
    with TestClient(app) as c:
        up = c.post(
            "/api/files",
            files={"file": ("sop-onboarding.txt", body, "text/plain")},
            data={"kind": "sop"},
            headers={"X-ERP-Delegation": token},
        )
        assert up.status_code == 201, up.text
        fid = up.json()["id"]
        bad = c.post(
            "/api/files",
            files={"file": ("x.pdf", b"not a pdf", "application/pdf")},
            data={"kind": "sop"},
            headers={"X-ERP-Delegation": token},
        )
        assert bad.status_code == 422
        drain()
        found = c.get("/api/files", params={"q": "laptop"}, headers={"X-ERP-Delegation": other}).json()["items"]
        assert found[0]["id"] == fid and "H-3" in found[0]["snippet"]
        got = c.get(f"/api/files/{fid}/content", headers={"X-ERP-Delegation": other})
        assert got.status_code == 200 and got.content == body
        with connect() as conn:
            assert one(conn, "SELECT action FROM file_access_log WHERE file_id=%s", (fid,))["action"] == "download"
        assert (
            c.post(f"/api/files/{fid}", json={"action": "withdraw"}, headers={"X-ERP-Delegation": other}).status_code
            == 403
        )
        assert c.get("/api/console/files", headers={"X-ERP-Delegation": token}).status_code == 401
        overview = c.get("/api/console/files", headers={"Authorization": "Bearer " + curator_token}).json()
        assert (
            any(r["origin"] == "managed" for r in overview["files"]) and overview["access"][0]["action"] == "download"
        )
        assert c.get("/api/files/kinds", headers={"X-ERP-Delegation": token}).json()["kinds"]
    assert hashlib  # sha256 of stored objects is verified on read (see files.content)
