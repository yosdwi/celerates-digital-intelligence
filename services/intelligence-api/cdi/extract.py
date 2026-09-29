"""One extractor for files (doc 17 §5): page-, table- and OCR-aware, shared by Company Files and Agent drops.

Output is a list of pages, each a list of blocks ``{"kind": "heading"|"text"|"table"|"ocr", "text": ...}``.

* Text PDFs take the fast path (pypdf, per page). A PDF with pages that have no text layer goes to Docling with OCR
  and table structure on, keeping page provenance (Docling's own models, pre-fetched into the image).
* DOCX keeps paragraphs and tables (as table blocks); XLSX/CSV become one page per sheet of table rows.
* Images of documents are OCR'd through Docling.

Nothing here decides access or classification. `pii_flags` only *detects* identity-document patterns so the registry
can hold a file for review; it never loosens anything."""

import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from tempfile import TemporaryDirectory
from xml.etree import ElementTree

MAX_PAGES = 300
MAX_OCR_PAGES = 50
MAX_CHARS = 600_000
CHUNK = 900
SPARSE = 5  # characters: a PDF page with less text than this has no usable text layer (a scan)
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

MEDIA = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".csv": "text/csv",
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
}


class ExtractError(ValueError):
    pass


@dataclass
class Extracted:
    pages: list = field(default_factory=list)  # [(page_no, [block, ...])]
    parser: str = ""
    ocr_pages: int = 0
    tables: int = 0
    scanned: bool = False  # content needs OCR that was not available

    @property
    def text(self):
        return "\n\n".join(b["text"] for _, blocks in self.pages for b in blocks)


def named(name, body):
    """`name` with an extension: ERP file columns carry labels such as "BAST", so infer it from the bytes."""
    if Path(name).suffix.lower() in MEDIA:
        return name
    head = body[:8]
    for magic, suffix in ((b"%PDF-", ".pdf"), (b"\x89PNG", ".png"), (b"\xff\xd8\xff", ".jpg")):
        if head.startswith(magic):
            return name + suffix
    return name


def sniff(name, body):
    """The media type from the bytes, checked against the extension. Uploaded names are not trusted."""
    suffix = Path(name).suffix.lower()
    if suffix not in MEDIA:
        raise ExtractError("Jenis berkas belum didukung.")
    head = body[:8]
    ok = {
        ".pdf": head.startswith(b"%PDF-"),
        ".docx": head.startswith(b"PK\x03\x04"),
        ".xlsx": head.startswith(b"PK\x03\x04"),
        ".png": head.startswith(b"\x89PNG\r\n\x1a\n"),
        ".jpg": head.startswith(b"\xff\xd8\xff"),
        ".jpeg": head.startswith(b"\xff\xd8\xff"),
    }.get(suffix, b"\x00" not in body[:4096])
    if not ok:
        raise ExtractError("Isi berkas tidak sesuai dengan jenisnya.")
    return suffix, MEDIA[suffix]


# ── per-format readers ─────────────────────────────────────────────────────────────────────────────────────────────
def pdf_pages(body):
    """[(page_no, text)] via pypdf. Raises ExtractError for unreadable or encrypted PDFs."""
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(body))
        if reader.is_encrypted:
            raise ExtractError("PDF terkunci kata sandi belum didukung.")
        return [(i + 1, page.extract_text() or "") for i, page in enumerate(reader.pages[:MAX_PAGES])]
    except ExtractError:
        raise
    except Exception as exc:
        raise ExtractError("PDF tidak dapat dibaca.") from exc


def docx_blocks(body):
    """Paragraphs and tables of a DOCX, in order. Tables become table blocks (rows as `a | b | c`)."""
    try:
        with zipfile.ZipFile(io.BytesIO(body)) as archive:
            info = archive.getinfo("word/document.xml")
            if info.file_size > 20 * 1024 * 1024:
                raise ExtractError("Isi DOCX terlalu besar.")
            root = ElementTree.fromstring(archive.read(info))
    except ExtractError:
        raise
    except Exception as exc:
        raise ExtractError("DOCX tidak dapat dibaca.") from exc
    blocks = []
    body_el = root.find(f"{W}body")
    for block in body_el if body_el is not None else []:
        if block.tag == f"{W}p":
            text = "".join(t.text or "" for t in block.iter(f"{W}t")).strip()
            style = block.find(f"{W}pPr/{W}pStyle")
            heading = style is not None and "heading" in (style.get(f"{W}val") or "").lower()
            if text:
                blocks.append({"kind": "heading" if heading else "text", "text": text})
        elif block.tag == f"{W}tbl":
            rows = []
            for row in block.iter(f"{W}tr"):
                cells = ["".join(t.text or "" for t in cell.iter(f"{W}t")).strip() for cell in row.iter(f"{W}tc")]
                rows.append(" | ".join(cells))
            if rows:
                blocks.append({"kind": "table", "text": "\n".join(rows)})
    return blocks


def sheet_pages(suffix, body, max_rows=2000):
    """One page per sheet; rows as table blocks with the header repeated every block (so each chunk is readable)."""
    sheets = []
    if suffix == ".csv":
        text = _decode(body)
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|") if text.strip() else csv.excel
        sheets.append(("CSV", list(csv.reader(io.StringIO(text), dialect))[:max_rows]))
    else:
        from openpyxl import load_workbook

        try:
            book = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
        except Exception as exc:
            raise ExtractError("XLSX tidak dapat dibaca.") from exc
        for ws in book.worksheets[:20]:
            rows = []
            for row in ws.iter_rows(values_only=True):
                cells = ["" if v is None else str(v).strip() for v in row]
                if any(cells):
                    rows.append(cells)
                if len(rows) >= max_rows:
                    break
            sheets.append((ws.title, rows))
    pages = []
    for no, (title, rows) in enumerate(sheets, 1):
        if not rows:
            continue
        header, blocks, group = " | ".join(rows[0]), [{"kind": "heading", "text": f"Sheet {title}"}], []
        for row in rows[1:]:
            group.append(" | ".join(row))
            if sum(len(r) for r in group) > CHUNK - len(header):
                blocks.append({"kind": "table", "text": header + "\n" + "\n".join(group)})
                group = []
        if group or len(rows) == 1:
            blocks.append({"kind": "table", "text": header + ("\n" + "\n".join(group) if group else "")})
        pages.append((no, blocks))
    return pages


def _decode(body):
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return body.decode(encoding)
        except UnicodeDecodeError:
            continue
    return ""


def _paragraph_blocks(text):
    blocks = []
    for para in re.split(r"\n\s*\n", text.replace("\r", "")):
        para = para.strip()
        if not para:
            continue
        line = " ".join(para.split())
        heading = (
            3 <= len(line) <= 80
            and not line.endswith((".", ",", ";"))
            and bool(re.match(r"^(\d+(\.\d+)*[.)]?|[A-Z]\.|BAB|Pasal|Section)\s", line) or line.isupper())
        )
        blocks.append({"kind": "heading" if heading else "text", "text": line if heading else para})
    return blocks


# ── Docling (layout, tables, OCR) ──────────────────────────────────────────────────────────────────────────────────
def docling_available():
    try:
        import docling  # noqa: F401

        return True
    except ImportError:
        return False


def _converter(ocr=True, tables=True):
    """Docling with RapidOCR on the torch backend (both already in the image) and TableFormer. Models come from
    DOCLING_ARTIFACTS_PATH when set (pre-fetched at image build), so the worker never downloads at run time."""
    from docling.datamodel.base_models import InputFormat
    from docling.datamodel.pipeline_options import PdfPipelineOptions, RapidOcrOptions
    from docling.document_converter import DocumentConverter, ImageFormatOption, PdfFormatOption

    options = PdfPipelineOptions(
        do_ocr=ocr, do_table_structure=tables, ocr_options=RapidOcrOptions(lang=["latin"], backend="torch")
    )
    return DocumentConverter(
        format_options={
            InputFormat.PDF: PdfFormatOption(pipeline_options=options),
            InputFormat.IMAGE: ImageFormatOption(pipeline_options=options),
        }
    )


def docling_pages(document):
    """A DoclingDocument → [(page_no, blocks)], keeping page provenance and tables. Separate from conversion so it is
    testable without models."""
    pages = {}
    tables = 0
    for item, _level in document.iterate_items():
        prov = getattr(item, "prov", None) or []
        page = prov[0].page_no if prov else 1
        label = str(getattr(item, "label", "")).lower()
        if label in ("page_header", "page_footer"):  # running headers/footers repeat on every page: noise
            continue
        if "table" in label and hasattr(item, "export_to_markdown"):
            text = item.export_to_markdown(doc=document) if _takes_doc(item) else item.export_to_markdown()
            kind = "table"
            tables += 1
        else:
            text = getattr(item, "text", "") or ""
            kind = "heading" if label in ("section_header", "title") else "text"
        if text.strip():
            pages.setdefault(page, []).append({"kind": kind, "text": text.strip()})
    return sorted(pages.items()), tables


def _takes_doc(item):
    import inspect

    try:
        return "doc" in inspect.signature(item.export_to_markdown).parameters
    except (TypeError, ValueError):
        return False


def _docling(suffix, body, ocr=True, tables=True, converter=None):
    with TemporaryDirectory() as directory:
        path = Path(directory) / ("source" + suffix)
        path.write_bytes(body)
        result = (converter or _converter(ocr=ocr, tables=tables)).convert(path)
    return docling_pages(result.document)


# ── entry point ────────────────────────────────────────────────────────────────────────────────────────────────────
def extract(name, body, *, want_tables=False, allow_ocr=True, converter=None):
    """Extract a file. `want_tables` sends text PDFs through Docling for table structure (contracts, BAST, manpower).
    `allow_ocr=False` (the synchronous Agent drop) marks a scanned file instead of running OCR."""
    suffix, _ = sniff(name, body)
    out = Extracted()
    if suffix in (".txt", ".md"):
        out.pages, out.parser = [(1, _paragraph_blocks(_decode(body)))], "text-v1"
    elif suffix in (".csv", ".xlsx"):
        out.pages, out.parser = sheet_pages(suffix, body), "sheets-v1"
        out.tables = sum(1 for _, bs in out.pages for b in bs if b["kind"] == "table")
    elif suffix == ".docx":
        out.pages, out.parser = [(1, docx_blocks(body))], "docx-v1"
        out.tables = sum(1 for b in out.pages[0][1] if b["kind"] == "table")
    elif suffix == ".pdf":
        raw = pdf_pages(body)
        sparse = [p for p, text in raw if len(text.strip()) < SPARSE]
        needs_docling = (sparse or want_tables) and (allow_ocr or not sparse)
        if needs_docling and (converter is not None or docling_available()):
            ocr = bool(sparse) and allow_ocr and len(sparse) <= MAX_OCR_PAGES
            pages, tables = _docling(suffix, body, ocr=ocr, tables=True, converter=converter)
            out.pages, out.parser, out.tables = pages, "docling-v2+ocr" if ocr else "docling-v2", tables
            out.ocr_pages = len(sparse) if ocr else 0
            if ocr:  # blocks on pages that had no text layer came from OCR
                out.pages = [
                    (p, [{**b, "kind": "ocr"} if p in sparse and b["kind"] == "text" else b for b in bs])
                    for p, bs in out.pages
                ]
        else:
            out.pages = [(p, _paragraph_blocks(text)) for p, text in raw if text.strip()]
            out.parser = "pypdf-v1"
        out.scanned = bool(sparse) and out.ocr_pages == 0 and not any(bs for _, bs in out.pages)
    else:  # image of a document
        if not allow_ocr or (converter is None and not docling_available()):
            out.scanned, out.parser = True, "none"
        else:
            out.pages, out.tables = _docling(suffix, body, ocr=True, tables=True, converter=converter)
            out.pages = [(p, [{**b, "kind": "ocr"} if b["kind"] == "text" else b for b in bs]) for p, bs in out.pages]
            out.parser, out.ocr_pages = "docling-v2+ocr", len(out.pages)
    return out


def chunks(pages, size=CHUNK):
    """Pack blocks into chunks of about `size` characters: never across pages, tables alone, headings carried."""
    out, total = [], 0
    for page, blocks in pages:
        heading, buffer, kind = None, "", "text"

        def flush():
            nonlocal buffer
            if buffer.strip():
                out.append({"ordinal": len(out) + 1, "page": page, "heading": heading, "block": kind, "text": buffer})
            buffer = ""

        for b in blocks:
            if b["kind"] == "heading":
                flush()
                heading = b["text"][:200]
                buffer = b["text"]
                kind = "text"
                continue
            if b["kind"] == "table":
                flush()
                kind = "table"
                buffer = b["text"][: size * 3]
                flush()
                kind = "text"
                continue
            text = " ".join(b["text"].split())
            if b["kind"] == "ocr":
                kind = "ocr"
            if buffer and len(buffer) + len(text) > size:
                flush()
            buffer = f"{buffer}\n{text}".strip() if buffer else text
            while len(buffer) > size * 2:
                cut = buffer.rfind(" ", 0, size) if " " in buffer[:size] else size
                out.append(
                    {"ordinal": len(out) + 1, "page": page, "heading": heading, "block": kind, "text": buffer[:cut]}
                )
                buffer = buffer[cut:].strip()
            total += len(text)
        flush()
        if total > MAX_CHARS:
            break
    return out


def headings(pages, limit=12):
    return list(dict.fromkeys(b["text"] for _, bs in pages for b in bs if b["kind"] == "heading"))[:limit]


# ── identity-document detection (tightening only) ─────────────────────────────────────────────────────────────────
PII = {
    "nik": re.compile(r"\b(?:NIK|N\.I\.K)\s*[:.]?\s*\d{16}\b|\b\d{16}\b"),
    "npwp": re.compile(r"\b\d{2}\.\d{3}\.\d{3}\.\d-\d{3}\.\d{3}\b"),
    "ktp": re.compile(r"KARTU\s+TANDA\s+PENDUDUK|KARTU\s+KELUARGA", re.I),
    "bank_account": re.compile(
        r"\b(?:no\.?\s*rek(?:ening)?|nomor\s+rekening|account\s+no\.?)\s*[:.]?\s*\d[\d\s-]{7,}", re.I
    ),
}


def pii_flags(text):
    """Names of identity-document patterns present. The matched values are never returned or stored."""
    return sorted(k for k, pattern in PII.items() if pattern.search(text or ""))
